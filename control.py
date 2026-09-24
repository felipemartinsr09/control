import os
import re
import glob
import argparse
from datetime import datetime
import pdfplumber
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Regras de categorias padrão para despesas
REGRAS_CATEGORIAS = {
    'Alimentação': [
        'carrefour', 'supermercado', 'mercado', 'feira', 'acougue', 'açougue',
        'padaria', 'panificadora', 'restaurante', 'ifood', 'burguer', 'burger',
        'mcdonalds', 'lanchonete', 'churrascaria', 'pizzaria', 'bh', 'extra',
        'assai', 'atacad', 'hortifruti', 'mercearia', 'cafe', 'boteco', 'bar'
    ],
    'Cuidados Pessoais': [
        'barber', 'barbearia', 'salao', 'salão', 'cabelo', 'estetica', 'beleza',
        'farmacia', 'farmácia', 'drogaria', 'pacheco', 'raia', 'drogasil',
        'venancio', 'cosmetico', 'perfumaria', 'corte'
    ],
    'Transporte': [
        'uber', '99', '99app', '99 *', 'posto', 'gasolina', 'combustivel',
        'shell', 'ipiranga', 'vibra', 'petrobras', 'estacionamento',
        'transcol', 'gvbus', 'pedagio', 'estac'
    ],
    'Lazer & Assinaturas': [
        'spotify', 'netflix', 'steam', 'cinema', 'prime', 'playstation',
        'game', 'disney', 'hbomax', 'sympla', 'ingresso'
    ],
    'Educação': [
        'faculdade', 'universidade', 'curso', 'livraria', 'udemy', 'escola'
    ],
    'Investimento': [
        'investimento', 'aplicacao', 'resgate', 'cdb', 'tesouro', 'b3', 'rico',
        'xp', 'nuinvest', 'inter dtvm', 'poupanca', 'rendimento', 'aporte'
    ],
    'Contas & Despesas Bancárias': [
        'tarifa', 'ted enviada', 'enel', 'cesan', 'edp', 'cobranca', 'debito automatico', 'iof'
    ]
}

# Paleta de cores pasteis para relatórios
CORES_CATEGORIAS = {
    'Recebimento': 'C6EFCE',                 # Verde pastel suave
    'Cartão': 'FADBD8',                      # Coral / Salmão suave
    'Transferência de Conta': 'D7CCC8',      # Bronze / neutro
    'Investimento': 'D9E1F2',                # Azul aço suave
    'Alimentação': 'FFE2D1',                 # Pêssego suave
    'Cuidados Pessoais': 'F8CECC',           # Rosa suave
    'Transporte': 'DAE8FC',                 # Azul claro
    'Lazer & Assinaturas': 'E1D5E7',         # Lilás pastel
    'Educação': 'D5E8D4',                    # Menta suave
    'Contas & Despesas Bancárias': 'FFF2CC', # Amarelo pastel
    'Outros': 'E0E0E0'                       # Cinza suave
}


def calcular_mes_posterior(data_str):
    """Calcula o mês posterior com tratamento de virada de ano."""
    partes = data_str.strip().split('/')
    mes = int(partes[1])
    ano = int(partes[2]) if len(partes) == 3 else datetime.now().year

    if mes == 12:
        mes_seguinte = 1
        ano_seguinte = ano + 1
    else:
        mes_seguinte = mes + 1
        ano_seguinte = ano

    return f"{mes_seguinte:02d}_{ano_seguinte}"


def classificar_transacao(texto, eh_credito):
    t = texto.lower()

    if eh_credito:
        return 'Recebimento'

    termos_cartao = ['fatura', 'pgto fatura', 'pagamento fatura', 'desconto cartao', 'desc cartao', 'pgto cartao']
    if any(tc in t for tc in termos_cartao):
        return 'Cartão'

    for termo in REGRAS_CATEGORIAS['Investimento']:
        if termo in t:
            return 'Investimento'

    for cat, termos in REGRAS_CATEGORIAS.items():
        if cat == 'Investimento':
            continue
        if any(termo in t for termo in termos):
            return cat

    return 'Outros'


def limpar_nome_estabelecimento(texto):
    t = texto
    lixos = [
        r'COMPRA CARTAO -', r'COMPRA CARTAO', r'COMPRA NACIONAL -', r'COMPRA NACIONAL',
        r'COMPRA ELO -', r'COMPRA MASTER -', r'COMPRA VISA -', r'COMPRA DEBITO -',
        r'PIX ENVIADO -', r'PIX TRANSF', r'PIX RECEBIDO -', r'PAGTO ELETRON COBRANCA',
        r'PAGTO TITULO', r'DEBITO CONTA CORRENTE', r'CREDITO EM CONTA',
        r'\bBRA\b', r'\bBR\b'
    ]
    for termo in lixos:
        t = re.sub(termo, '', t, flags=re.IGNORECASE)

    t = re.sub(r'\b\d{4,}\b', '', t)
    t = re.sub(r'\s+', ' ', t).strip(' -/*')
    return t.title() if len(t) > 1 else 'Transação Diversa'


def processar_extrato(caminho_pdf):
    transacoes = []
    re_data = re.compile(r'^(\d{2}/\d{2}(?:/\d{4})?)')
    re_valor = re.compile(r'(\d{1,3}(?:\.\d{3})*,\d{2})\s*([DC-])?')
    ignorar = ['SALDO ANTERIOR', 'SALDO ATUAL', 'SALDO DO DIA', 'S A L D O', 'TOTAL:', 'EXTRATO']

    try:
        with pdfplumber.open(caminho_pdf) as pdf:
            todas_linhas = []
            for pagina in pdf.pages:
                texto = pagina.extract_text()
                if texto:
                    todas_linhas.extend([l.strip() for l in texto.split('\n') if l.strip()])
    except Exception as e:
        print(f"❌ Erro ao abrir o arquivo PDF: {e}")
        return

    i = 0
    while i < len(todas_linhas):
        linha = todas_linhas[i]
        
        if any(termo in linha.upper() for termo in ignorar):
            i += 1
            continue

        match_data = re_data.search(linha)
        matches_valor = list(re_valor.finditer(linha))

        if match_data and matches_valor:
            data = match_data.group(1)
            ultimo_valor = matches_valor[-1]
            valor_str = ultimo_valor.group(1).replace('.', '').replace(',', '.')
            indicador = (ultimo_valor.group(2) or '').upper()

            eh_debito = 'D' in indicador or '-' in linha or linha.endswith('D')
            eh_credito = 'C' in indicador or ('C' in linha.split()[-1] and not eh_debito)

            try:
                valor = abs(float(valor_str))
            except ValueError:
                i += 1
                continue

            pos_ini = match_data.end()
            pos_fim = ultimo_valor.start()
            desc_linha = linha[pos_ini:pos_fim].strip()

            desc_completa = desc_linha
            if i + 1 < len(todas_linhas):
                proxima_linha = todas_linhas[i + 1]
                if not re_data.search(proxima_linha) and not re_valor.search(proxima_linha):
                    desc_completa = f"{desc_linha} {proxima_linha}".strip()
                    i += 1

            categoria = classificar_transacao(desc_completa, eh_credito)
            nome_final = limpar_nome_estabelecimento(desc_completa)
            tipo_fluxo = 'Entrada' if eh_credito else 'Saída'

            transacoes.append({
                'Data': data,
                'Estabelecimento / Descrição': nome_final,
                'Tipo': tipo_fluxo,
                'Categoria': categoria,
                'Valor (R$)': valor
            })

        i += 1

    if not transacoes:
        print("❌ Nenhuma transação identificada no arquivo.")
        return

    df = pd.DataFrame(transacoes)

    # Identificação do mês posterior
    primeira_data = df['Data'].iloc[0]
    mes_posterior_ano = calcular_mes_posterior(primeira_data)
    nome_arquivo = f"Gastos_{mes_posterior_ano}.xlsx"

    with pd.ExcelWriter(nome_arquivo, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Extrato Geral', index=False)

    wb = load_workbook(nome_arquivo)
    ws_geral = wb['Extrato Geral']

    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

    for col_num in range(1, 6):
        c = ws_geral.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center')

    for row in range(2, len(df) + 2):
        cat_linha = ws_geral.cell(row=row, column=4).value
        cor_hex = CORES_CATEGORIAS.get(cat_linha, 'FFFFFF')
        ws_geral.cell(row=row, column=4).fill = PatternFill(start_color=cor_hex, end_color=cor_hex, fill_type='solid')
        ws_geral.cell(row=row, column=5).number_format = 'R$ #,##0.00'

    total_recebimentos = df[df['Tipo'] == 'Entrada']['Valor (R$)'].sum()
    total_gastos = df[df['Tipo'] == 'Saída']['Valor (R$)'].sum()
    saldo_liquido = total_recebimentos - total_gastos

    ws_painel = wb.create_sheet(title='Painel por Categoria')

    borda_fina = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    # Card 1: Total Recebimentos
    ws_painel.merge_cells('A1:B1')
    ws_painel['A1'] = "TOTAL RECEBIMENTOS"
    ws_painel['A1'].font = Font(bold=True, color='276A3C', size=11)
    ws_painel['A1'].alignment = Alignment(horizontal='center')
    ws_painel['A1'].fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

    ws_painel.merge_cells('A2:B2')
    ws_painel['A2'] = total_recebimentos
    ws_painel['A2'].font = Font(bold=True, size=13)
    ws_painel['A2'].number_format = 'R$ #,##0.00'
    ws_painel['A2'].alignment = Alignment(horizontal='center')

    # Card 2: Total Gastos
    ws_painel.merge_cells('D1:E1')
    ws_painel['D1'] = "TOTAL GASTOS (DESPESAS)"
    ws_painel['D1'].font = Font(bold=True, color='9C0006', size=11)
    ws_painel['D1'].alignment = Alignment(horizontal='center')
    ws_painel['D1'].fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

    ws_painel.merge_cells('D2:E2')
    ws_painel['D2'] = total_gastos
    ws_painel['D2'].font = Font(bold=True, size=13)
    ws_painel['D2'].number_format = 'R$ #,##0.00'
    ws_painel['D2'].alignment = Alignment(horizontal='center')

    # Card 3: Saldo Líquido
    ws_painel.merge_cells('G1:H1')
    ws_painel['G1'] = "SALDO LÍQUIDO DO MÊS"
    ws_painel['G1'].font = Font(bold=True, color='1F497D', size=11)
    ws_painel['G1'].alignment = Alignment(horizontal='center')
    cor_saldo = 'D9E1F2' if saldo_liquido >= 0 else 'F8CECC'
    ws_painel['G1'].fill = PatternFill(start_color=cor_saldo, end_color=cor_saldo, fill_type='solid')

    ws_painel.merge_cells('G2:H2')
    ws_painel['G2'] = saldo_liquido
    ws_painel['G2'].font = Font(bold=True, size=13)
    ws_painel['G2'].number_format = 'R$ #,##0.00'
    ws_painel['G2'].alignment = Alignment(horizontal='center')

    col_offset = 1
    ordem_categorias = [
        'Recebimento', 'Cartão', 'Transferência de Conta', 'Investimento',
        'Alimentação', 'Cuidados Pessoais', 'Transporte',
        'Lazer & Assinaturas', 'Educação', 'Contas & Despesas Bancárias', 'Outros'
    ]

    for cat in ordem_categorias:
        df_cat = df[df['Categoria'] == cat]
        if df_cat.empty:
            continue

        cor_hex = CORES_CATEGORIAS.get(cat, 'E0E0E0')
        cat_fill = PatternFill(start_color=cor_hex, end_color=cor_hex, fill_type='solid')

        c_tit = ws_painel.cell(row=5, column=col_offset, value=cat.upper())
        ws_painel.merge_cells(start_row=5, start_column=col_offset, end_row=5, end_column=col_offset + 1)
        c_tit.fill = cat_fill
        c_tit.font = Font(bold=True, size=11)
        c_tit.alignment = Alignment(horizontal='center')

        soma_cat = df_cat['Valor (R$)'].sum()
        c_soma = ws_painel.cell(row=6, column=col_offset, value=f"Soma: R$ {soma_cat:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
        ws_painel.merge_cells(start_row=6, start_column=col_offset, end_row=6, end_column=col_offset + 1)
        c_soma.font = Font(bold=True, italic=True)
        c_soma.alignment = Alignment(horizontal='center')

        ws_painel.cell(row=7, column=col_offset, value="Origem / Destino").font = Font(bold=True)
        ws_painel.cell(row=7, column=col_offset + 1, value="Valor").font = Font(bold=True)

        row_item = 8
        for _, reg in df_cat.iterrows():
            c1 = ws_painel.cell(row=row_item, column=col_offset, value=reg['Estabelecimento / Descrição'])
            c2 = ws_painel.cell(row=row_item, column=col_offset + 1, value=reg['Valor (R$)'])
            c2.number_format = 'R$ #,##0.00'
            c1.border = borda_fina
            c2.border = borda_fina
            row_item += 1

        col_offset += 3

    for ws in [ws_geral, ws_painel]:
        for col in ws.columns:
            valores = [str(cell.value or '') for cell in col]
            max_len = max([len(v) for v in valores] + [10])
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

    wb.save(nome_arquivo)

    print("\n" + "=" * 55)
    print(f"✅ EXCEL CONSOLIDADO GERADO: {nome_arquivo}")
    print(f"🟢 Total Recebimentos: R$ {total_recebimentos:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
    print(f"🔴 Total Gastos:       R$ {total_gastos:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
    print(f"⚖️ Saldo Líquido:      R$ {saldo_liquido:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
    print("=" * 55 + "\n")


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Control - Processador de extratos bancários em PDF para planilhas organizadas.")
    parser.add_argument('arquivo', nargs='?', help='Caminho do arquivo PDF do extrato (opcional).')
    args = parser.parse_args()

    if args.arquivo:
        if os.path.exists(args.arquivo):
            print(f"📄 Processando: {args.arquivo}")
            processar_extrato(args.arquivo)
        else:
            print(f"❌ Arquivo não encontrado: {args.arquivo}")
    else:
        arquivos_pdf = glob.glob("*.pdf")
        if not arquivos_pdf:
            print("❌ Nenhum extrato em PDF encontrado na pasta.")
        else:
            pdf_recente = max(arquivos_pdf, key=os.path.getctime)
            print(f"📄 Processando mais recente: {pdf_recente}")
            processar_extrato(pdf_recente)